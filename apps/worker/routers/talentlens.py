import os
import json
import logging
import asyncio
import threading
import zipfile
import tempfile
import io
from typing import List, Dict, Any, Optional
from fastapi import APIRouter, UploadFile, File, Form, HTTPException, Depends, Request
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
import pdfplumber
from sentence_transformers import SentenceTransformer
from groq import AsyncGroq, RateLimitError

logger = logging.getLogger("hirevault-talentlens")

router = APIRouter()

embedding_model = None
model_lock = threading.Lock()

def get_embedding_model():
    global embedding_model
    with model_lock:
        if embedding_model is None:
            logger.info("Loading sentence-transformers model...")
            embedding_model = SentenceTransformer('all-MiniLM-L6-v2')
    return embedding_model

def get_1536_embedding(text: str) -> List[float]:
    try:
        model = get_embedding_model()
        emb = model.encode(text).tolist()
        emb.extend([0.0] * (1536 - len(emb)))
        return emb
    except Exception as e:
        logger.error(f"SentenceTransformer failed on this system: {e}. Falling back to zero-vector.")
        return [0.0] * 1536

def get_supabase():
    supabase_url = os.getenv("NEXT_PUBLIC_SUPABASE_URL")
    supabase_key = os.getenv("SUPABASE_SERVICE_ROLE_KEY")
    if not supabase_url or not supabase_key:
        raise HTTPException(status_code=500, detail="Supabase env vars missing")
    from supabase import create_client
    return create_client(supabase_url, supabase_key)

def verify_hiring_manager(request: Request):
    user_id = request.headers.get("x-user-id")
    if not user_id:
        raise HTTPException(status_code=401, detail="Unauthorized")
    return user_id

PARSE_SYSTEM_PROMPT = """You are an expert technical recruiter and resume parser.
Extract information from the provided CV text.
Return ONLY valid JSON. No markdown, no preamble.
Schema:
{
  "parsed": {
    "name": "string",
    "email": "string",
    "skills": ["string"],
    "roles": ["string"],
    "years_experience": number,
    "companies": ["string"],
    "summary": "string"
  },
  "scores": {
    "skills": number (0-100),
    "seniority": number (0-100),
    "trajectory": number (0-100),
    "culture": number (0-100)
  },
  "archetype": "Perfect fit" | "High ceiling" | "Solid hire" | "Overqualified" | "Needs review"
}
"""

async def process_single_cv(filename: str, raw_text: str, job_posting_id: str, seq_idx: int, sem: asyncio.Semaphore):
    async with sem:
        groq = AsyncGroq(api_key=os.getenv("GROQ_API_KEY"))
        model = os.getenv("GROQ_MODEL", "llama-3.3-70b-versatile")
        
        try:
            models_to_try = [
                os.getenv("GROQ_MODEL", "llama-3.3-70b-versatile"),
                "llama-3.1-8b-instant",
                "mixtral-8x7b-32768",
                "gemma2-9b-it"
            ]
            resp = None
            for idx, m in enumerate(models_to_try):
                try:
                    resp = await groq.chat.completions.create(
                        messages=[
                            {"role": "system", "content": PARSE_SYSTEM_PROMPT},
                            {"role": "user", "content": f"CV Text:\n{raw_text[:6000]}"}
                        ],
                        model=m,
                        temperature=0.1,
                        response_format={"type": "json_object"}
                    )
                    break
                except Exception as e:
                    if "rate_limit_exceeded" in str(e).lower() or "429" in str(e) or isinstance(e, RateLimitError):
                        if idx < len(models_to_try) - 1:
                            logger.warning(f"Rate limit on {m}. Falling back to {models_to_try[idx+1]}...")
                            continue
                    raise e
                    
            parsed_data = json.loads(resp.choices[0].message.content)
            
            emb = await asyncio.to_thread(get_1536_embedding, raw_text)
            
            scores = parsed_data.get("scores", {})
            comp_score = (scores.get("skills", 0) * 0.40) + \
                         (scores.get("seniority", 0) * 0.25) + \
                         (scores.get("trajectory", 0) * 0.20) + \
                         (scores.get("culture", 0) * 0.15)
                         
            anonymized_id = f"Candidate #{seq_idx:03d}"
            
            return {
                "job_posting_id": job_posting_id,
                "filename": filename,
                "raw_text": raw_text,
                "parsed_json": parsed_data.get("parsed", {}),
                "embedding": emb,
                "composite_score": round(comp_score, 1),
                "skills_score": scores.get("skills", 0),
                "seniority_score": scores.get("seniority", 0),
                "trajectory_score": scores.get("trajectory", 0),
                "culture_score": scores.get("culture", 0),
                "archetype": parsed_data.get("archetype", "Needs review"),
                "status": "new",
                "anonymized_id": anonymized_id
            }
        except Exception as e:
            logger.error(f"Failed to process {filename}: {e}")
            return None

def extract_text_from_file(filename: str, content: bytes) -> str:
    filename = filename.lower()
    text = ""
    try:
        if filename.endswith(".pdf"):
            with pdfplumber.open(io.BytesIO(content)) as pdf:
                text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        elif filename.endswith(".docx"):
            from docx import Document
            doc = Document(io.BytesIO(content))
            text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
        else:
            try:
                text = content.decode("utf-8", errors="ignore")
            except Exception:
                pass
    except Exception as e:
        logger.warning(f"Failed to extract text from {filename}: {e}")
    return text.strip()

@router.post("/upload-bulk")
async def upload_bulk(
    request: Request,
    job_posting_id: str = Form(...),
    files: List[UploadFile] = File(...),
):
    hm_id = verify_hiring_manager(request)
    supabase = get_supabase()
    
    cv_texts = []
    
    logger.info(f"Received {len(files)} files for bulk upload")
    for file in files:
        logger.info(f"Incoming file name: {file.filename}, content-type: {file.content_type}")
        content = await file.read()
        if file.filename.lower().endswith(".zip"):
            with zipfile.ZipFile(io.BytesIO(content)) as z:
                for zinfo in z.infolist():
                    if "__MACOSX" not in zinfo.filename and not zinfo.is_dir():
                        with z.open(zinfo) as pf:
                            file_content = pf.read()
                            text = extract_text_from_file(zinfo.filename, file_content)
                            if text:
                                cv_texts.append((zinfo.filename, text))
        else:
            text = extract_text_from_file(file.filename, content)
            if text:
                cv_texts.append((file.filename, text))
                    
    if not cv_texts:
        raise HTTPException(status_code=400, detail="No readable documents found.")

    sem = asyncio.Semaphore(10)
    tasks = []
    
    res = await asyncio.to_thread(lambda: supabase.table("cv_submissions").select("id", count="exact").eq("job_posting_id", job_posting_id).execute())
    current_count = res.count if getattr(res, 'count', None) is not None else 0
    
    for idx, (fname, text) in enumerate(cv_texts):
        tasks.append(process_single_cv(fname, text, job_posting_id, current_count + idx + 1, sem))
        
    results = await asyncio.gather(*tasks)
    valid_results = [r for r in results if r is not None]
    
    if not valid_results and cv_texts:
        raise HTTPException(status_code=500, detail="Failed to process any CVs. This is likely due to exceeding your daily Groq API limits.")
    
    if valid_results:
        await asyncio.to_thread(lambda: supabase.table("cv_submissions").insert(valid_results).execute())
        
    return {"processed": len(valid_results), "job_posting_id": job_posting_id}

class RerankRequest(BaseModel):
    job_posting_id: str
    jd_text: str
    weights: Dict[str, float]
    top_n: int

@router.post("/rerank")
async def rerank(request: Request, req: RerankRequest):
    hm_id = verify_hiring_manager(request)
    supabase = get_supabase()
    
    res = await asyncio.to_thread(lambda: supabase.table("cv_submissions").select("*").eq("job_posting_id", req.job_posting_id).execute())
    cvs = res.data
    
    if not cvs:
        return []
        
    # Optional: We could use jd embedding to refine scores, but for now we just apply new weights
    # JD embedding is computed as requested, though logic implies we apply new weights to existing dimension scores.
    jd_emb = await asyncio.to_thread(get_1536_embedding, req.jd_text)
    
    updates = []
    w_s = req.weights.get("skills", 0) / 100.0
    w_sen = req.weights.get("seniority", 0) / 100.0
    w_t = req.weights.get("trajectory", 0) / 100.0
    w_c = req.weights.get("culture", 0) / 100.0
    
    for cv in cvs:
        new_comp = (cv["skills_score"] * w_s) + \
                   (cv["seniority_score"] * w_sen) + \
                   (cv["trajectory_score"] * w_t) + \
                   (cv["culture_score"] * w_c)
        cv["composite_score"] = round(new_comp, 1)
        updates.append({"id": cv["id"], "composite_score": cv["composite_score"]})
        
    # Bulk update in supabase is tricky without RPC, so we do it one by one or via upsert
    if updates:
        await asyncio.to_thread(lambda: supabase.table("cv_submissions").upsert(updates).execute())
        
    cvs.sort(key=lambda x: x["composite_score"], reverse=True)
    return cvs[:req.top_n]

class ChatCVRequest(BaseModel):
    cv_submission_id: str
    question: str
    conversation_history: List[Dict[str, str]]

@router.post("/chat-cv")
async def chat_cv(request: Request, req: ChatCVRequest):
    hm_id = verify_hiring_manager(request)
    supabase = get_supabase()
    
    res = await asyncio.to_thread(lambda: supabase.table("cv_submissions").select("parsed_json, raw_text, job_postings!inner(hiring_manager_id)").eq("id", req.cv_submission_id).execute())
    if not res.data or res.data[0]["job_postings"]["hiring_manager_id"] != hm_id:
        raise HTTPException(status_code=403, detail="Forbidden")
        
    raw_text = res.data[0]["raw_text"] or ""
    parsed_json = res.data[0].get("parsed_json")
    
    cv_context = json.dumps(parsed_json, indent=2) if parsed_json else (raw_text[:8000] if len(raw_text) > 8000 else raw_text)
    
    async def stream_generator():
        groq = AsyncGroq(api_key=os.getenv("GROQ_API_KEY"))
        messages = [
            {"role": "system", "content": f"You are a recruiting assistant. Answer questions about the following CV only. Be concise.\nIf the answer isn't present in the provided CV text, say so explicitly rather than guessing.\nCV:\n{cv_context}"}
        ]
        messages.extend(req.conversation_history)
        messages.append({"role": "user", "content": req.question})
        
        stream = await groq.chat.completions.create(
            messages=messages,
            model=os.getenv("GROQ_MODEL", "llama-3.3-70b-versatile"),
            temperature=0.3,
            stream=True
        )
        async for chunk in stream:
            if chunk.choices[0].delta.content is not None:
                yield chunk.choices[0].delta.content
                
    return StreamingResponse(stream_generator(), media_type="text/event-stream")

class EmailRequest(BaseModel):
    cv_submission_id: str
    tone: str
    anonymized: bool

@router.post("/generate-email")
async def generate_email(request: Request, req: EmailRequest):
    hm_id = verify_hiring_manager(request)
    supabase = get_supabase()
    
    res = await asyncio.to_thread(lambda: supabase.table("cv_submissions").select("parsed_json, job_postings!inner(hiring_manager_id)").eq("id", req.cv_submission_id).execute())
    if not res.data or res.data[0]["job_postings"]["hiring_manager_id"] != hm_id:
        raise HTTPException(status_code=403, detail="Forbidden")
        
    parsed = res.data[0].get("parsed_json", {})
    
    if req.anonymized:
        return {
            "subject": "Discussion regarding your application",
            "body": "Hi [NAME],\n\nI was impressed by your experience at [COMPANY] and your skills in [SKILLS]. We would love to discuss a potential fit for our open role.\n\nBest,\nHiring Team"
        }
        
    groq = AsyncGroq(api_key=os.getenv("GROQ_API_KEY"))
    prompt = f"""You are a recruiter writing a {req.tone} outreach email.
Candidate Name: {parsed.get('name', '[Name]')}
Top Skills: {', '.join(parsed.get('skills', [])[:3])}
Recent Company: {', '.join(parsed.get('companies', [])[:1])}
Summary: {parsed.get('summary', '')}

Only reference candidate details explicitly provided above. Do not invent past employers, achievements, or skills not listed. Use [TOKEN] placeholders (e.g. [CALENDAR_LINK], [ROLE_TITLE]) for any detail the recruiter must fill in themselves.
Return JSON format: {{"subject": "...", "body": "..."}}"""

    resp = await groq.chat.completions.create(
        messages=[{"role": "user", "content": prompt}],
        model=os.getenv("GROQ_MODEL", "llama-3.3-70b-versatile"),
        temperature=0.4,
        response_format={"type": "json_object"}
    )
    
    return json.loads(resp.choices[0].message.content)
